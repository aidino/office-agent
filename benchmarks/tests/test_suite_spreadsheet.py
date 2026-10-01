"""Tests for the SpreadsheetBench 2 suite adapter.

Tests monkeypatch ``_recalculate`` to identity (fixtures carry literal cell
values, not formulas).  The recalc-failure path is tested separately by
monkeypatching ``_recalculate`` to return ``None``.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from office_bench.suites.spreadsheet import SpreadsheetSuite
from office_bench.suites.base import AgentOutput

FIXTURES = Path(__file__).parent / "fixtures" / "spreadsheet"


@pytest.fixture(autouse=True)
def _no_recalc(monkeypatch) -> None:
    """Skip LibreOffice recalc in unit tests — fixtures write literal values."""
    monkeypatch.setattr(
        SpreadsheetSuite, "_recalculate", staticmethod(lambda p: p)
    )


@pytest.fixture()
def suite() -> SpreadsheetSuite:
    return SpreadsheetSuite(FIXTURES)


def _make_output() -> AgentOutput:
    return AgentOutput(
        messages=[{"role": "assistant", "content": "done"}],
        files_created=[],
        tool_calls=[],
        duration_seconds=4.0,
    )


# ── load_tasks ──────────────────────────────────────────────────────


def test_load_tasks(suite: SpreadsheetSuite) -> None:
    tasks = suite.load_tasks()
    # 3 tasks across Debugging, Financial_Model, Visualization
    assert len(tasks) == 3
    categories = {t.category for t in tasks}
    assert "Debugging" in categories
    assert "Financial_Model" in categories


def test_task_fields(suite: SpreadsheetSuite) -> None:
    tasks = suite.load_tasks()
    t = next(t for t in tasks if t.task_id == "debug-001")
    assert t.suite == "spreadsheet"
    assert t.category == "Debugging"
    assert "SUM" in t.prompt


# ── format_prompt ───────────────────────────────────────────────────


def test_format_prompt(suite: SpreadsheetSuite) -> None:
    tasks = suite.load_tasks()
    t = tasks[0]
    prompt = suite.format_prompt(t)
    assert len(prompt) > 0


# ── evaluate: deterministic cell comparison ─────────────────────────


def test_evaluate_cell_value_pass(suite: SpreadsheetSuite, tmp_path: Path) -> None:
    tasks = suite.load_tasks()
    t = next(t for t in tasks if t.task_id == "debug-001")

    # Create output xlsx with correct cell value
    wb = Workbook()
    ws = wb.active
    ws["B5"] = 150
    wb.save(tmp_path / "debug-001.xlsx")

    result = suite.evaluate(t, tmp_path, _make_output())
    assert result.passed is True
    assert result.score == 1.0
    assert result.judge_backend == "deterministic"


def test_evaluate_cell_value_fail(suite: SpreadsheetSuite, tmp_path: Path) -> None:
    tasks = suite.load_tasks()
    t = next(t for t in tasks if t.task_id == "debug-001")

    wb = Workbook()
    ws = wb.active
    ws["B5"] = 999  # wrong value
    wb.save(tmp_path / "debug-001.xlsx")

    result = suite.evaluate(t, tmp_path, _make_output())
    assert result.passed is False
    assert result.score < 1.0


def test_evaluate_multiple_cells(suite: SpreadsheetSuite, tmp_path: Path) -> None:
    tasks = suite.load_tasks()
    t = next(t for t in tasks if t.task_id == "fin-001")

    wb = Workbook()
    ws = wb.active
    ws["C3"] = 1000
    ws["C4"] = 1100
    wb.save(tmp_path / "fin-001.xlsx")

    result = suite.evaluate(t, tmp_path, _make_output())
    assert result.passed is True
    assert result.score == 1.0


def test_evaluate_partial_cell_match(suite: SpreadsheetSuite, tmp_path: Path) -> None:
    """One cell correct, one wrong → partial score, not passed."""
    tasks = suite.load_tasks()
    t = next(t for t in tasks if t.task_id == "fin-001")

    wb = Workbook()
    ws = wb.active
    ws["C3"] = 1000  # correct
    ws["C4"] = 9999  # wrong
    wb.save(tmp_path / "fin-001.xlsx")

    result = suite.evaluate(t, tmp_path, _make_output())
    assert result.passed is False
    assert result.score == 0.5  # 1/2 cells matched


# ── evaluate: Visualization deferred ────────────────────────────────




# ── load_tasks edge cases ──────────────────────────────────────────


def test_load_tasks_empty_repo(tmp_path: Path) -> None:
    """Repo with no data/ dir returns empty list."""
    suite = SpreadsheetSuite(tmp_path)
    assert suite.load_tasks() == []


def test_load_tasks_skips_non_dirs(tmp_path: Path) -> None:
    """Files in data/ dir are skipped."""
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    (data_dir / "not_a_dir.txt").write_text("junk")
    suite = SpreadsheetSuite(tmp_path)
    assert suite.load_tasks() == []


def test_load_tasks_skips_missing_dataset_json(tmp_path: Path) -> None:
    """Category dir without dataset.json is skipped."""
    cat_dir = tmp_path / "data" / "SomeCategory"
    cat_dir.mkdir(parents=True)
    suite = SpreadsheetSuite(tmp_path)
    assert suite.load_tasks() == []


def test_load_tasks_skips_bad_json(tmp_path: Path) -> None:
    """Malformed dataset.json is skipped with a warning."""
    cat_dir = tmp_path / "data" / "BadCat"
    cat_dir.mkdir(parents=True)
    (cat_dir / "dataset.json").write_text("{broken")
    suite = SpreadsheetSuite(tmp_path)
    assert suite.load_tasks() == []


def test_load_tasks_skips_entry_without_task_id(tmp_path: Path) -> None:
    """Entry with empty task_id is skipped."""
    import json as _json

    cat_dir = tmp_path / "data" / "Debugging"
    cat_dir.mkdir(parents=True)
    (cat_dir / "dataset.json").write_text(
        _json.dumps([{"instruction": "no id", "spreadsheet_file": "x.xlsx"}])
    )
    suite = SpreadsheetSuite(tmp_path)
    assert suite.load_tasks() == []


# ── setup_workspace ────────────────────────────────────────────────


def test_setup_workspace_copies_files(suite: SpreadsheetSuite, tmp_path: Path) -> None:
    """Input files are copied into workspace."""
    import json as _json

    # Create a dummy spreadsheet in the fixture tree
    cat_dir = FIXTURES / "data" / "Debugging"
    xlsx_path = cat_dir / "debug-001.xlsx"
    created = False
    if not xlsx_path.exists():
        wb = Workbook()
        wb.save(xlsx_path)
        created = True

    try:
        tasks = suite.load_tasks()
        t = next(t for t in tasks if t.task_id == "debug-001")
        suite.setup_workspace(t, tmp_path)
        # If the file existed as input, it should be copied
        if t.input_files:
            for f in t.input_files:
                assert (tmp_path / f.name).exists()
    finally:
        if created:
            xlsx_path.unlink(missing_ok=True)


# ── _compare_cell edge cases ──────────────────────────────────────


def test_compare_cell_none_vs_empty() -> None:
    """None actual matches empty expected string."""
    assert SpreadsheetSuite._compare_cell(None, "", 0.0) is True


def test_compare_cell_none_vs_nonempty() -> None:
    """None actual does not match non-empty expected."""
    assert SpreadsheetSuite._compare_cell(None, "100", 0.0) is False


def test_compare_cell_string_fallback() -> None:
    """Non-numeric strings use exact string comparison."""
    assert SpreadsheetSuite._compare_cell("hello", "hello", 0.0) is True
    assert SpreadsheetSuite._compare_cell("hello", "world", 0.0) is False


# ── evaluate: no expected cells (empty check) ──────────────────────


def test_evaluate_no_expected_cells(suite: SpreadsheetSuite, tmp_path: Path) -> None:
    """Task with no expected_cells passes with score 1.0."""
    tasks = suite.load_tasks()
    t = next(t for t in tasks if t.task_id == "viz-001")
    # Override to a deterministic category with no cells
    from dataclasses import replace
    t2 = replace(t, category="Template", metadata={"expected_cells": {}, "spreadsheet_file": "x.xlsx"})
    (tmp_path / "x.xlsx").write_bytes(b"fake")
    result = suite.evaluate(t2, tmp_path, _make_output())
    assert result.passed is True
    assert result.score == 1.0
def test_evaluate_visualization_deferred(suite: SpreadsheetSuite, tmp_path: Path) -> None:
    tasks = suite.load_tasks()
    t = next(t for t in tasks if t.task_id == "viz-001")
    result = suite.evaluate(t, tmp_path, _make_output())
    assert result.score == 0.0
    assert "deferred" in result.notes.lower() or "visualization" in result.notes.lower()


# ── evaluate: error paths ───────────────────────────────────────────


def test_evaluate_missing_file(suite: SpreadsheetSuite, tmp_path: Path) -> None:
    tasks = suite.load_tasks()
    t = next(t for t in tasks if t.task_id == "debug-001")
    result = suite.evaluate(t, tmp_path, _make_output())
    assert result.passed is False


def test_evaluate_recalc_failure_fails_cleanly(
    suite: SpreadsheetSuite, tmp_path: Path, monkeypatch
) -> None:
    """soffice missing/crashing must fail the task with an explicit note."""
    tasks = suite.load_tasks()
    t = next(t for t in tasks if t.task_id == "debug-001")

    wb = Workbook()
    ws = wb.active
    ws["B5"] = 150
    wb.save(tmp_path / "debug-001.xlsx")

    monkeypatch.setattr(
        SpreadsheetSuite, "_recalculate", staticmethod(lambda p: None)
    )
    result = suite.evaluate(t, tmp_path, _make_output())
    assert result.passed is False
    assert "recalc" in result.notes.lower()


# ── Suite protocol ──────────────────────────────────────────────────


def test_suite_satisfies_protocol(suite: SpreadsheetSuite) -> None:
    from office_bench.suites.base import Suite
    assert isinstance(suite, Suite)
