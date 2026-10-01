"""Tests for the PPTC suite adapter."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

from office_bench.suites.base import AgentOutput, Suite
from office_bench.suites.pptc import PPTCSuite

FIXTURES = Path(__file__).parent / "fixtures" / "pptc"


@pytest.fixture(autouse=True)
def _create_fixture_pptx_files() -> None:
    """Create template.pptx (edit fixture) and the session_1 label pptx."""
    template_path = FIXTURES / "PPT_test_input" / "Edit_ppt_template" / "template.pptx"
    if not template_path.exists():
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "Original Title"
        template_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(template_path))

    # Label deck for session_1 (what main.py --prepare generates):
    # slide 1 title "Hello World"; slide 2 title + two bullet points.
    label_path = FIXTURES / "PPT_label_Create_new_slides" / "session_1.pptx"
    if not label_path.exists():
        prs = Presentation()
        s1 = prs.slides.add_slide(prs.slide_layouts[0])
        s1.shapes.title.text = "Hello World"
        s2 = prs.slides.add_slide(prs.slide_layouts[1])
        s2.shapes.title.text = "Agenda"
        tf = s2.placeholders[1].text_frame  # type: ignore[index]
        tf.text = "First point"
        p = tf.add_paragraph()
        p.text = "Second point"
        label_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(label_path))


@pytest.fixture()
def suite() -> PPTCSuite:
    return PPTCSuite(FIXTURES)


def _make_output() -> AgentOutput:
    return AgentOutput(
        messages=[{"role": "assistant", "content": "done"}],
        files_created=[],
        tool_calls=[],
        duration_seconds=5.0,
    )


# ── load_tasks ──────────────────────────────────────────────────────


def test_load_tasks(suite: PPTCSuite) -> None:
    tasks = suite.load_tasks()
    assert len(tasks) == 2
    ids = {t.task_id for t in tasks}
    assert "session_1" in ids
    assert "session_2" in ids


def test_task_fields(suite: PPTCSuite) -> None:
    tasks = suite.load_tasks()
    t1 = next(t for t in tasks if t.task_id == "session_1")
    assert t1.suite == "pptc"
    assert t1.category == "Create_new_slides"
    assert len(t1.metadata.get("turns", [])) == 2
    assert t1.metadata["turn_prompts"] == [
        "Create a title slide with text 'Hello World'",
        "Add a second slide with bullet points",
    ]
    assert t1.metadata["label_file"] is not None


def test_load_tasks_empty_dir(tmp_path: Path) -> None:
    """Suite pointed at a dir with no PPT_test_input/ returns []."""
    s = PPTCSuite(tmp_path)
    assert s.load_tasks() == []


# ── setup_workspace ─────────────────────────────────────────────────


def test_setup_workspace_copies_template(suite: PPTCSuite, tmp_path: Path) -> None:
    tasks = suite.load_tasks()
    t2 = next(t for t in tasks if t.task_id == "session_2")
    suite.setup_workspace(t2, tmp_path)
    assert (tmp_path / "template.pptx").exists()


def test_setup_workspace_noop_for_create_task(suite: PPTCSuite, tmp_path: Path) -> None:
    """Create-type tasks have no template — workspace stays empty."""
    tasks = suite.load_tasks()
    t1 = next(t for t in tasks if t.task_id == "session_1")
    suite.setup_workspace(t1, tmp_path)
    assert list(tmp_path.iterdir()) == []


# ── format_prompt ───────────────────────────────────────────────────


def test_format_prompt_first_turn(suite: PPTCSuite) -> None:
    tasks = suite.load_tasks()
    t1 = next(t for t in tasks if t.task_id == "session_1")
    prompt = suite.format_prompt(t1)
    assert "Hello World" in prompt


# ── evaluate (PPTX-Match) ──────────────────────────────────────────


def _save_two_slide_prediction(tmp_path: Path, title: str) -> None:
    """Prediction deck with the same shape as the session_1 label."""
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = title
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Agenda"
    tf = s2.placeholders[1].text_frame  # type: ignore[index]
    tf.text = "First point"
    p = tf.add_paragraph()
    p.text = "Second point"
    prs.save(str(tmp_path / "prediction.pptx"))


def test_evaluate_match_passes(suite: PPTCSuite, tmp_path: Path) -> None:
    tasks = suite.load_tasks()
    t1 = next(t for t in tasks if t.task_id == "session_1")

    _save_two_slide_prediction(tmp_path, title="Hello World")
    result = suite.evaluate(t1, tmp_path, _make_output())

    assert result.suite == "pptc"
    assert result.task_id == "session_1"
    assert result.judge_backend == "deterministic"
    assert result.passed is True
    assert result.score == 1.0


def test_evaluate_wrong_content_fails(suite: PPTCSuite, tmp_path: Path) -> None:
    """A deck with any-text-but-wrong-content must NOT pass."""
    tasks = suite.load_tasks()
    t1 = next(t for t in tasks if t.task_id == "session_1")

    _save_two_slide_prediction(tmp_path, title="Wrong Title")
    result = suite.evaluate(t1, tmp_path, _make_output())

    assert result.passed is False
    assert result.score < 1.0


def test_evaluate_missing_prediction_fails(suite: PPTCSuite, tmp_path: Path) -> None:
    """No pptx in workspace → fail with explicit note."""
    tasks = suite.load_tasks()
    t1 = next(t for t in tasks if t.task_id == "session_1")

    result = suite.evaluate(t1, tmp_path, _make_output())
    assert result.passed is False
    assert result.score == 0.0
    assert "prediction" in result.notes.lower() or "pptx" in result.notes.lower()


def test_evaluate_missing_label_reports_prepare_step(
    suite: PPTCSuite, tmp_path: Path
) -> None:
    """session_2 has no label file — must fail with a prepare hint."""
    tasks = suite.load_tasks()
    t2 = next(t for t in tasks if t.task_id == "session_2")

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0])
    prs.save(str(tmp_path / "template.pptx"))

    result = suite.evaluate(t2, tmp_path, _make_output())
    assert result.passed is False
    assert result.score == 0.0
    assert "prepare" in result.notes.lower()


def test_evaluate_extra_slides_fails(suite: PPTCSuite, tmp_path: Path) -> None:
    """Prediction with extra slides beyond label should fail."""
    tasks = suite.load_tasks()
    t1 = next(t for t in tasks if t.task_id == "session_1")

    # Create a 3-slide prediction (label has 2)
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Hello World"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Agenda"
    tf = s2.placeholders[1].text_frame  # type: ignore[index]
    tf.text = "First point"
    p = tf.add_paragraph()
    p.text = "Second point"
    s3 = prs.slides.add_slide(prs.slide_layouts[0])
    s3.shapes.title.text = "Extra slide"
    prs.save(str(tmp_path / "prediction.pptx"))

    result = suite.evaluate(t1, tmp_path, _make_output())
    assert result.passed is False


# ── protocol conformance ────────────────────────────────────────────


def test_suite_satisfies_protocol(suite: PPTCSuite) -> None:
    """PPTCSuite is a structural subtype of Suite."""
    assert isinstance(suite, Suite)
