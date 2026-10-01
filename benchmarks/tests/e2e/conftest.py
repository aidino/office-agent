"""Shared fixtures for E2E tests.

These tests exercise the full CLI → Runner → real suite adapter → evaluation
pipeline. The only mock is AgentBridge (no live gateway), but everything
else — suite loading, workspace setup, evaluation logic, result persistence,
backdata CSV, and Markdown report — runs for real.
"""

from __future__ import annotations

import json
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from openpyxl import Workbook
from pptx import Presentation

from office_bench.suites.base import AgentOutput
from office_bench.suites.forte import ForteSuite
from office_bench.suites.officebench import OfficeBenchSuite
from office_bench.suites.pptc import PPTCSuite
from office_bench.suites.spreadsheet import SpreadsheetSuite

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures"


@pytest.fixture()
def results_dir(tmp_path: Path) -> Path:
    """Isolated results directory for one test."""
    d = tmp_path / "results" / "benchmarks"
    d.mkdir(parents=True)
    return d


@pytest.fixture()
def mock_bridge() -> MagicMock:
    """Bridge mock whose run_task creates files the evaluators expect."""
    bridge = MagicMock()

    def _run_task(prompt: str, workspace_dir: Path) -> AgentOutput:
        """Produce output files matching what each suite's evaluator needs."""
        # OfficeBench: evaluate_file_exist checks budget.xlsx
        if "budget" in prompt.lower():
            wb = Workbook()
            wb.save(workspace_dir / "budget.xlsx")

        # OfficeBench: evaluate_contain checks summary.txt contains "total"
        if "summarize" in prompt.lower():
            (workspace_dir / "summary.txt").write_text(
                "The total revenue is $500K."
            )

        # SpreadsheetBench: cell-level comparison
        if "fix the sum" in prompt.lower() or "fix" in prompt.lower():
            wb = Workbook()
            ws = wb.active
            ws["B5"] = 150
            wb.save(workspace_dir / "debug-001.xlsx")

        if "revenue projection" in prompt.lower():
            wb = Workbook()
            ws = wb.active
            ws["C3"] = 1000
            ws["C4"] = 1100
            wb.save(workspace_dir / "fin-001.xlsx")

        return AgentOutput(
            messages=[{
                "role": "assistant",
                "content": (
                    "Done. Total revenue is $500K. "
                    "Department A employees listed."
                ),
            }],
            files_created=[],
            tool_calls=[{"name": "write_file", "args": {}}],
            duration_seconds=5.0,
        )

    def _run_session(prompts: list[str], workspace_dir: Path) -> AgentOutput:
        """Produce a pptx for PPTC multi-turn sessions."""
        prs = Presentation()
        for prompt in prompts:
            slide_layout = prs.slide_layouts[0]
            slide = prs.slides.add_slide(slide_layout)
            if slide.shapes.title:
                slide.shapes.title.text = (
                    prompt[:40] if len(prompt) > 40 else prompt
                )
        prs.save(workspace_dir / "output.pptx")

        return AgentOutput(
            messages=[{
                "role": "assistant",
                "content": f"Created presentation with {len(prompts)} slides.",
            }],
            files_created=[Path("output.pptx")],
            tool_calls=[],
            duration_seconds=3.0 * len(prompts),
        )

    bridge.run_task.side_effect = _run_task
    bridge.run_session.side_effect = _run_session
    return bridge


@pytest.fixture(autouse=True)
def _ensure_forte_assets() -> None:
    """Create fixture xlsx so FORTE setup_workspace can copy it."""
    assets_dir = FIXTURES / "forte" / "data" / "assets" / "finance-001" / "input"
    assets_dir.mkdir(parents=True, exist_ok=True)
    data_file = assets_dir / "data.xlsx"
    if not data_file.exists():
        data_file.write_bytes(b"fake xlsx content")


@pytest.fixture(autouse=True)
def _ensure_spreadsheet_files() -> None:
    """Create fixture xlsx files SpreadsheetSuite.setup_workspace needs."""
    for category in ("Debugging", "Financial_Model"):
        cat_dir = FIXTURES / "spreadsheet" / "data" / category
        dataset = cat_dir / "dataset.json"
        if not dataset.exists():
            continue
        entries = json.loads(dataset.read_text())
        for entry in entries:
            xlsx_file = cat_dir / entry.get("spreadsheet_file", "")
            if xlsx_file.name and not xlsx_file.exists():
                wb = Workbook()
                wb.save(xlsx_file)


@pytest.fixture(autouse=True)
def _ensure_pptc_template() -> None:
    """Create template pptx for PPTC Edit tasks."""
    template_dir = (
        FIXTURES / "pptc" / "PPT_test_input" / "Edit_ppt_template"
    )
    template_file = template_dir / "template.pptx"
    if not template_file.exists():
        prs = Presentation()
        slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(slide_layout)
        if slide.shapes.title:
            slide.shapes.title.text = "Original Title"
        prs.save(template_file)


@pytest.fixture()
def patched_cli(results_dir: Path, mock_bridge: MagicMock):
    """Patch CLI globals so runs write to tmp_path and use mock bridge.

    We patch _build_suite_registry to return suite adapters pointing at
    the test fixtures directory rather than DATA_BASE (which points at
    data/benchmarks/ submodules that may not be cloned).
    """
    def _fixture_registry() -> dict:
        return {
            "officebench": OfficeBenchSuite(FIXTURES / "officebench"),
            "pptc": PPTCSuite(FIXTURES / "pptc"),
            "spreadsheet": SpreadsheetSuite(FIXTURES / "spreadsheet"),
            "forte": ForteSuite(FIXTURES / "forte"),
        }

    with (
        patch("office_bench.cli.RESULTS_BASE", results_dir),
        patch("office_bench.cli._build_suite_registry", _fixture_registry),
        patch("office_bench.runner.AgentBridge", return_value=mock_bridge),
    ):
        yield results_dir
