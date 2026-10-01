"""Fixture stand-in for OfficeBench's native evaluation module.

Mirrors the upstream API surface used by the adapter. After `setup` clones
the real submodule, verify the real function names/signatures (Step 5
investigation) and keep this fixture in sync.
"""

from __future__ import annotations

from pathlib import Path


def evaluate_file_exist(args: dict, workspace_dir: Path) -> bool:
    return (workspace_dir / args.get("file_path", "")).exists()


def evaluate_contain(args: dict, workspace_dir: Path) -> bool:
    file_path = workspace_dir / args.get("file_path", "")
    if not file_path.exists():
        return False
    content = file_path.read_text(errors="replace")
    return str(args.get("expected", "")).lower() in content.lower()


def evaluate_exact_match(args: dict, workspace_dir: Path) -> bool:
    file_path = workspace_dir / args.get("file_path", "")
    if not file_path.exists():
        return False
    content = file_path.read_text(errors="replace").strip()
    return content == str(args.get("expected", "")).strip()


def evaluate_excel_cell_value(args: dict, workspace_dir: Path) -> bool:
    from openpyxl import load_workbook

    file_path = workspace_dir / args.get("file_path", "")
    if not file_path.exists():
        return False
    wb = load_workbook(file_path, data_only=True)
    ws = wb[args["sheet"]] if args.get("sheet") else wb.active
    actual = ws[args.get("cell", "A1")].value
    return actual is not None and str(actual).strip() == str(
        args.get("expected", "")
    ).strip()
