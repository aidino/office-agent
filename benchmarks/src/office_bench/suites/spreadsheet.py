"""SpreadsheetBench 2 suite adapter — deterministic cell-level comparison.

Supports three deterministic categories (Debugging, Financial_Model, Template).
Visualization is deferred to Phase 6 (requires VLM).

Before cell comparison, ``evaluate()`` recalculates the workbook via
LibreOffice headless — openpyxl-written files carry no cached formula
results, so ``data_only=True`` would read ``None`` for every formula cell.
Unit tests monkeypatch ``_recalculate`` to identity because fixtures write
literal values.
"""

from __future__ import annotations

import json
import logging
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from office_bench.suites.base import AgentOutput, Task, TaskResult

_log = logging.getLogger(__name__)

_DETERMINISTIC_CATEGORIES = frozenset({"Debugging", "Financial_Model", "Template"})


class SpreadsheetSuite:
    """Adapter for SpreadsheetBench 2 (RUCKBReasoning/SpreadsheetBench-2)."""

    name = "spreadsheet"

    def __init__(self, repo_dir: Path) -> None:
        self._repo_dir = repo_dir

    # ── Suite protocol ──────────────────────────────────────────────

    def load_tasks(self) -> list[Task]:
        """Parse ``data/<category>/dataset.json`` for each category."""
        data_dir = self._repo_dir / "data"
        if not data_dir.exists():
            return []

        tasks: list[Task] = []
        for cat_dir in sorted(data_dir.iterdir()):
            if not cat_dir.is_dir():
                continue
            dataset_file = cat_dir / "dataset.json"
            if not dataset_file.exists():
                continue

            try:
                entries = json.loads(dataset_file.read_text())
            except (json.JSONDecodeError, OSError):
                _log.warning("Skipping unreadable dataset: %s", dataset_file)
                continue

            for entry in entries:
                task = self._parse_entry(entry, cat_dir)
                if task is not None:
                    tasks.append(task)
        return tasks

    def setup_workspace(self, task: Task, workspace_dir: Path) -> None:
        """Copy spreadsheet files into workspace."""
        for src in task.input_files:
            if src.is_dir():
                shutil.copytree(src, workspace_dir / src.name, dirs_exist_ok=True)
            else:
                shutil.copy2(src, workspace_dir / src.name)

    def format_prompt(self, task: Task) -> str:
        """Return the task instruction."""
        return task.prompt

    def evaluate(
        self,
        task: Task,
        workspace_dir: Path,
        agent_output: AgentOutput,
    ) -> TaskResult:
        """Cell-level comparison for deterministic categories; defer Visualization."""
        category = task.category

        if category == "Visualization":
            return TaskResult(
                task_id=task.task_id,
                suite=self.name,
                passed=False,
                score=0.0,
                breakdown={},
                notes="Visualization evaluation deferred (requires VLM, Phase 6)",
                judge_backend=None,
            )

        expected_cells: dict = task.metadata.get("expected_cells", {})
        spreadsheet_file: str = task.metadata.get("spreadsheet_file", "")
        output_path = workspace_dir / spreadsheet_file

        if not output_path.exists():
            return TaskResult(
                task_id=task.task_id,
                suite=self.name,
                passed=False,
                score=0.0,
                breakdown={},
                notes=f"Output file not found: {spreadsheet_file}",
                judge_backend="deterministic",
            )

        if not expected_cells:
            return TaskResult(
                task_id=task.task_id,
                suite=self.name,
                passed=True,
                score=1.0,
                breakdown={},
                notes="No cells to check",
                judge_backend="deterministic",
            )

        # LibreOffice recalc (spec §5.3 prerequisite)
        recalced = self._recalculate(output_path)
        if recalced is None:
            return TaskResult(
                task_id=task.task_id,
                suite=self.name,
                passed=False,
                score=0.0,
                breakdown={},
                notes=(
                    "LibreOffice recalc unavailable or failed — formula "
                    "results cannot be compared"
                ),
                judge_backend="deterministic",
            )

        try:
            wb = load_workbook(recalced, data_only=True)
            ws = wb.active
        except Exception as exc:
            return TaskResult(
                task_id=task.task_id,
                suite=self.name,
                passed=False,
                score=0.0,
                breakdown={},
                notes=f"Error loading workbook: {exc}",
                judge_backend="deterministic",
            )

        matches = 0
        total = len(expected_cells)
        breakdown: dict[str, float] = {}

        for cell_ref, expected_info in expected_cells.items():
            expected_value = str(expected_info.get("value", ""))
            tolerance = float(expected_info.get("tolerance", 0.0))
            actual_value = ws[cell_ref].value

            cell_match = self._compare_cell(actual_value, expected_value, tolerance)
            breakdown[cell_ref] = 1.0 if cell_match else 0.0
            if cell_match:
                matches += 1

        score = matches / total if total > 0 else 0.0
        passed = matches == total

        return TaskResult(
            task_id=task.task_id,
            suite=self.name,
            passed=passed,
            score=score,
            breakdown=breakdown,
            notes=f"{matches}/{total} cells matched",
            judge_backend="deterministic",
        )

    # ── internals ───────────────────────────────────────────────────

    def _parse_entry(self, entry: dict, cat_dir: Path) -> Task | None:
        """Parse a single dataset entry into a Task."""
        task_id = entry.get("task_id", "")
        if not task_id:
            _log.warning("Skipping entry without task_id in %s", cat_dir)
            return None

        spreadsheet_file = entry.get("spreadsheet_file", "")
        input_files: list[Path] = []
        if spreadsheet_file:
            sf = cat_dir / spreadsheet_file
            if sf.exists():
                input_files.append(sf)

        return Task(
            suite=self.name,
            task_id=task_id,
            prompt=entry.get("instruction", ""),
            category=entry.get("category", cat_dir.name),
            input_files=input_files,
            metadata={
                "expected_cells": entry.get("expected_cells", {}),
                "spreadsheet_file": spreadsheet_file,
                "category_dir": str(cat_dir),
            },
        )

    @staticmethod
    def _compare_cell(
        actual: object, expected_str: str, tolerance: float
    ) -> bool:
        """Compare a cell value against expected, with numeric tolerance."""
        if actual is None:
            return expected_str == ""

        actual_str = str(actual).strip()

        # Try numeric comparison with tolerance
        try:
            actual_num = float(actual_str)
            expected_num = float(expected_str)
            return abs(actual_num - expected_num) <= tolerance + abs(expected_num) * tolerance
        except (ValueError, TypeError):
            pass

        # Fall back to string comparison
        return actual_str == expected_str.strip()

    _SOFFICE = "soffice"

    @classmethod
    def _recalculate(cls, xlsx_path: Path) -> Path | None:
        """Recalculate formulas via LibreOffice headless.

        Returns the recalculated copy's path, or ``None`` on failure.
        Unit tests monkeypatch this method to identity — fixtures write
        literal values, so no recalc is needed there.
        """
        try:
            with tempfile.TemporaryDirectory() as tmp:
                subprocess.run(
                    [
                        cls._SOFFICE,
                        "--headless",
                        "--convert-to",
                        "xlsx",
                        "--outdir",
                        tmp,
                        str(xlsx_path),
                    ],
                    capture_output=True,
                    timeout=120,
                    check=True,
                )
                converted = Path(tmp) / xlsx_path.name
                if not converted.exists():
                    return None
                target = xlsx_path.with_suffix(".recalced.xlsx")
                shutil.copy2(converted, target)
                return target
        except (OSError, subprocess.SubprocessError):
            return None
