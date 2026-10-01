from __future__ import annotations

import pytest
from openpyxl import load_workbook
from pypdf import PdfWriter
from pptx import Presentation

from office_agent import tools


@pytest.fixture()
def workspace(tmp_path, monkeypatch):
    monkeypatch.setenv(tools.WORKSPACE_ENV, str(tmp_path))
    return tmp_path


def test_read_text_file_returns_content(workspace, tmp_path) -> None:
    (tmp_path / "notes.txt").write_text("quarterly report draft", encoding="utf-8")

    assert tools.read_text_file("notes.txt") == "quarterly report draft"


def test_read_text_file_accepts_absolute_path_inside_workspace(workspace, tmp_path) -> None:
    target = tmp_path / "nested" / "data.csv"
    target.parent.mkdir()
    target.write_text("a,b\n1,2\n", encoding="utf-8")

    content = tools.read_text_file(str(target))

    assert "1,2" in content


def test_read_text_file_rejects_path_escape(workspace, tmp_path) -> None:
    outside = tmp_path.parent / "secret.txt"
    outside.write_text("top secret", encoding="utf-8")

    result = tools.read_text_file(str(outside))

    assert result.startswith("Error: ")
    assert "workspace" in result


def test_read_text_file_rejects_relative_escape() -> None:
    assert tools.read_text_file("../../etc/passwd").startswith("Error: ")


def test_read_text_file_reports_missing_file(workspace) -> None:
    assert tools.read_text_file("missing.txt").startswith("Error: ")


def test_read_pdf_file_counts_pages(workspace, tmp_path) -> None:
    pdf_path = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with pdf_path.open("wb") as handle:
        writer.write(handle)

    result = tools.read_pdf_file("blank.pdf")

    assert "1 page(s)" in result
    assert "--- page 1 ---" in result


def test_spreadsheet_roundtrip(workspace, tmp_path) -> None:
    wrote = tools.write_spreadsheet(
        "reports/sales.xlsx",
        sheet_name="Q1",
        rows=[["item", "units"], ["pen", 10], ["paper", 25]],
    )

    assert wrote.startswith("Wrote 3 row(s)")

    result = tools.read_spreadsheet("reports/sales.xlsx")
    assert "Sheet 'Q1' (3 rows x 2 cols)" in result
    assert "| pen | 10 |" in result

    workbook = load_workbook(tmp_path / "reports" / "sales.xlsx")
    assert workbook.active.title == "Q1"


def test_write_presentation_creates_deck(workspace, tmp_path) -> None:
    wrote = tools.write_presentation(
        "reports/summary.pptx",
        title="Q1 Review",
        slides=[
            {"title": "Highlights", "bullets": ["Revenue up", "", "Churn flat"]},
            {"title": "Next steps", "bullets": ["Hire two engineers"]},
        ],
    )

    assert wrote.startswith("Wrote 2 slide(s)")

    deck = Presentation(tmp_path / "reports" / "summary.pptx")
    titles = [slide.shapes.title.text for slide in deck.slides]
    assert titles == ["Q1 Review", "Highlights", "Next steps"]


def test_write_tools_reject_path_escape(workspace) -> None:
    escape_rows = tools.write_spreadsheet("../outside.xlsx", "S", [[1]])
    escape_deck = tools.write_presentation("../outside.pptx", "T", [])

    assert escape_rows.startswith("Error: ")
    assert escape_deck.startswith("Error: ")
